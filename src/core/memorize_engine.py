import json
import os
import random
from datetime import date, timedelta

from src.utils import get_resource_path, get_user_data_path

# 艾宾浩斯遗忘曲线复习间隔（学习后第N天）：1天后、2天后、4天后、7天后、15天后
REVIEW_INTERVALS = [1, 2, 4, 7, 15]

WORDS_PER_CARD = 6
CARDS_PER_PAGE = 2
CARD_EN_FONT = 16      # 英文/词组基准字号
CARD_CN_FONT = 13.5    # 中文释义基准字号


class MemorizeEngine:
    def __init__(self, banks_file=None, progress_file=None):
        if banks_file is None:
            banks_file = get_resource_path(os.path.join("src", "data", "memorize_banks.json"))
        if progress_file is None:
            progress_file = get_user_data_path("memorize_progress.json")
        self.progress_file = progress_file
        self.days = self._load_banks(banks_file)
        self.progress = self._load_progress()

    # ---------- 数据加载 ----------

    def _load_banks(self, path):
        if not os.path.exists(path):
            return []
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data.get("days", [])

    def _load_progress(self):
        if os.path.exists(self.progress_file):
            try:
                with open(self.progress_file, "r", encoding="utf-8") as f:
                    return json.load(f)
            except (json.JSONDecodeError, IOError):
                pass
        return {"modules": {}}

    def save_progress(self):
        with open(self.progress_file, "w", encoding="utf-8") as f:
            json.dump(self.progress, f, indent=4, ensure_ascii=False)

    # ---------- 模块查询 ----------

    def get_days(self):
        return self.days

    def get_day(self, day_no):
        for d in self.days:
            if d["day"] == day_no:
                return d
        return None

    def get_day_entries(self, day_no):
        """展平某天的全部词条：[{text, pos, translation, category}]"""
        day = self.get_day(day_no)
        if not day:
            return []
        entries = []
        for cat in day["categories"]:
            for w in cat["words"]:
                entries.append({
                    "text": w.get("text", ""),
                    "pos": w.get("pos", ""),
                    "translation": w.get("translation", ""),
                    "category": cat["name"],
                })
        return entries

    def to_practice_list(self, day_no):
        """转成口语练习可用的词条列表（发音跟读用）"""
        items = []
        for e in self.get_day_entries(day_no):
            items.append({
                "text": e["text"],
                "translation": e["translation"],
                "phonetic": "",
                "group": f"速记Day{day_no}",
            })
        return items

    # ---------- 学习进度与复习计划 ----------

    def start_module(self, day_no, start_date=None):
        """开始学习某天模块，记录起始日期并生成复习计划"""
        key = str(day_no)
        modules = self.progress["modules"]
        if key not in modules:
            if start_date is None:
                start_date = date.today().isoformat()
            modules[key] = {
                "start_date": start_date,
                "stage": 0,
                "last_review": "",
                "tests": [],
            }
            self.save_progress()

    def is_started(self, day_no):
        return str(day_no) in self.progress["modules"]

    def get_module_progress(self, day_no):
        """返回模块进度信息；未开始返回 None"""
        m = self.progress["modules"].get(str(day_no))
        if not m:
            return None
        start = date.fromisoformat(m["start_date"])
        stage = m["stage"]
        total = len(REVIEW_INTERVALS)
        if stage >= total:
            next_due = None
        else:
            next_due = (start + timedelta(days=REVIEW_INTERVALS[stage])).isoformat()
        last_test = m["tests"][-1] if m["tests"] else None
        return {
            "stage": stage,
            "total_stages": total,
            "start_date": m["start_date"],
            "next_due": next_due,
            "last_review": m.get("last_review", ""),
            "last_test": last_test,
        }

    def get_due_reviews(self, today=None):
        """获取今天（含逾期）应复习的模块列表"""
        if today is None:
            today = date.today()
        due = []
        for key, m in self.progress["modules"].items():
            stage = m["stage"]
            if stage >= len(REVIEW_INTERVALS):
                continue
            start = date.fromisoformat(m["start_date"])
            due_date = start + timedelta(days=REVIEW_INTERVALS[stage])
            if due_date <= today:
                due.append({
                    "day": int(key),
                    "stage": stage,
                    "due_date": due_date.isoformat(),
                    "overdue_days": (today - due_date).days,
                })
        due.sort(key=lambda x: x["due_date"])
        return due

    def complete_review(self, day_no):
        """完成一次复习，推进阶段"""
        m = self.progress["modules"].get(str(day_no))
        if not m:
            return False
        m["stage"] = min(m["stage"] + 1, len(REVIEW_INTERVALS))
        m["last_review"] = date.today().isoformat()
        self.save_progress()
        return True

    def record_test(self, day_no, score, total):
        m = self.progress["modules"].get(str(day_no))
        if not m:
            return
        m["tests"].append({
            "date": date.today().isoformat(),
            "score": score,
            "total": total,
        })
        self.save_progress()

    def reset_all_progress(self):
        self.progress = {"modules": {}}
        self.save_progress()

    # ---------- 记忆卡片 ----------

    def build_cards(self, day_no, per_card=WORDS_PER_CARD):
        """切分卡片。末尾只剩 1~3 个词条时并入前面 1~2 张卡，不单独成卡。"""
        entries = self.get_day_entries(day_no)
        if not entries:
            return []
        cards = [entries[i:i + per_card] for i in range(0, len(entries), per_card)]
        if len(cards) >= 2 and 1 <= len(cards[-1]) <= 3:
            leftover = cards.pop()
            half = len(leftover) // 2
            if half:
                cards[-2].extend(leftover[:half])
            cards[-1].extend(leftover[half:])
        return cards

    def export_cards_pdf(self, day_no, out_path):
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.units import mm
            from reportlab.pdfgen import canvas
        except ImportError:
            raise RuntimeError("缺少 reportlab 库，请先安装：pip install reportlab")

        font, bold_font = _ensure_pdf_fonts()
        cards = self.build_cards(day_no)
        day = self.get_day(day_no)
        title = day["title"] if day else ""
        total = len(cards)

        pw, ph = A4
        margin = 12 * mm
        cut_gap = 5 * mm
        card_h = (ph - 2 * margin - cut_gap) / 2
        card_w = pw - 2 * margin

        c = canvas.Canvas(out_path, pagesize=A4)
        c.setTitle(f"单词速记卡 Day {day_no}")

        all_entries = [e for c_ in cards for e in c_]
        x_idx = margin + 5 * mm
        x_en = margin + 12 * mm
        right = margin + card_w - 5 * mm
        gap = 1 * mm
        pos_w = _pos_col_width([e["pos"] for e in all_entries], font, 10 * MM_TO_PT)
        en_w, cn_w = _split_card_columns(
            [e["text"] for e in all_entries],
            [e["translation"] for e in all_entries],
            bold_font, font, CARD_EN_FONT, CARD_CN_FONT,
            right - x_en - pos_w - gap)
        x_pos = x_en + en_w + gap
        x_cn = x_pos + pos_w

        for idx, card in enumerate(cards):
            pos_on_page = idx % CARDS_PER_PAGE
            if idx > 0 and pos_on_page == 0:
                c.showPage()
            top = ph - margin - pos_on_page * (card_h + cut_gap)
            bottom = top - card_h

            # 卡片外框（仅作裁剪参考，卡内无表格线）
            c.setStrokeColorRGB(0.75, 0.75, 0.75)
            c.setLineWidth(0.7)
            c.roundRect(margin, bottom, card_w, card_h, 4 * mm)
            if pos_on_page == 0 and idx < total - 1:
                c.setDash(3, 3)
                c.line(margin, bottom - cut_gap / 2,
                       margin + card_w, bottom - cut_gap / 2)
                c.setDash()

            # 顶部居中标题
            cx = margin + card_w / 2
            header_h = 16 * mm
            c.setFillColorRGB(0.05, 0.25, 0.55)
            c.setFont(bold_font, 17)
            c.drawCentredString(cx, top - 7.5 * mm,
                                f"单词速记卡（第{day_no}天：{idx + 1:03d}/{total:03d}）")
            if title:
                c.setFillColorRGB(0.55, 0.55, 0.55)
                c.setFont(font, 10)
                c.drawCentredString(cx, top - 13.2 * mm, title)

            # 词条区：按卡片行数平分整块高度，字号随之放大
            area_top = top - header_h
            area_bottom = bottom + 4 * mm
            row_h = (area_top - area_bottom) / len(card)

            en_size = _pick_card_font_size([e["text"] for e in card], bold_font,
                                           en_w, CARD_EN_FONT, row_h)
            cn_size = _pick_card_font_size([e["translation"] for e in card], font,
                                           cn_w, CARD_CN_FONT, row_h)

            for r, entry in enumerate(card):
                center_y = area_top - (r + 0.5) * row_h

                c.setFillColorRGB(0.7, 0.7, 0.7)
                c.setFont(font, 10)
                c.drawString(x_idx, center_y - 3.5, str(r + 1))

                if entry["pos"]:
                    c.setFillColorRGB(0.5, 0.5, 0.5)
                    c.setFont(font, 10)
                    c.drawString(x_pos, center_y - 3.5, entry["pos"])

                _draw_lines(c, bold_font, entry["text"], x_en, en_w, center_y, en_size,
                            (0.1, 0.1, 0.1))
                _draw_lines(c, font, entry["translation"], x_cn, cn_w, center_y, cn_size,
                            (0.08, 0.35, 0.7))

        c.save()
        return out_path

    def export_cards_docx(self, day_no, out_path):
        try:
            from docx import Document
            from docx.shared import Mm, Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
            from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
        except ImportError:
            raise RuntimeError("缺少 python-docx 库，请先安装：pip install python-docx")

        cards = self.build_cards(day_no)
        day = self.get_day(day_no)
        title = day["title"] if day else ""
        total = len(cards)

        doc = Document()
        section = doc.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(12)
        section.bottom_margin = Mm(12)
        section.left_margin = Mm(14)
        section.right_margin = Mm(14)

        content_w = 182.0          # 210 - 14 - 14
        card_h = (273.0 - 4.0) / 2 - 1.5   # 半页高度（含少量余量，避免溢出到下一页）
        header_h = 15.0
        cell_pad = 4.0             # 单元格左右内边距合计（mm）

        def set_run(run, size, bold=False, color=None):
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.name = "微软雅黑"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
            if color:
                run.font.color.rgb = RGBColor(*color)

        def clean_para(paragraph, align=None):
            fmt = paragraph.paragraph_format
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(0)
            fmt.line_spacing = 1.0
            if align:
                paragraph.alignment = align

        def borderless(table):
            tbl_pr = table._tbl.tblPr
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                el = OxmlElement(f"w:{edge}")
                el.set(qn("w:val"), "none")
                el.set(qn("w:sz"), "0")
                borders.append(el)
            tbl_pr.append(borders)
            layout = OxmlElement("w:tblLayout")
            layout.set(qn("w:type"), "fixed")
            tbl_pr.append(layout)

        def spacer(height, dashed_line=False, page_break=False):
            p = doc.add_paragraph()
            clean_para(p)
            p.paragraph_format.line_spacing = height
            run = p.add_run()
            run.font.size = Pt(2)
            if dashed_line:
                p_pr = p._p.get_or_add_pPr()
                borders = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "dashed")
                bottom.set(qn("w:sz"), "6")
                bottom.set(qn("w:space"), "1")
                bottom.set(qn("w:color"), "AAAAAA")
                borders.append(bottom)
                p_pr.append(borders)
            if page_break:
                run.add_break(WD_BREAK.PAGE)
            return p

        font, bold_font = _ensure_pdf_fonts()

        all_en = [e["text"] for c in cards for e in c]
        all_cn = [e["translation"] for c in cards for e in c]
        idx_w = 8.0
        pos_w = _pos_col_width([e["pos"] for c in cards for e in c], font,
                               10 * MM_TO_PT) / MM_TO_PT
        avail_mm = content_w - idx_w - pos_w - 2 * cell_pad
        en_pt, cn_pt = _split_card_columns(
            all_en, all_cn, bold_font, font, CARD_EN_FONT, CARD_CN_FONT,
            avail_mm * MM_TO_PT)
        col_w = [Mm(idx_w), Mm(en_pt / MM_TO_PT + cell_pad), Mm(pos_w),
                 Mm(cn_pt / MM_TO_PT + cell_pad)]
        en_w_pt = en_pt
        cn_w_pt = cn_pt

        for idx, card in enumerate(cards):
            table = doc.add_table(rows=len(card) + 1, cols=4)
            table.autofit = False
            borderless(table)

            hdr = table.rows[0]
            hdr.height = Mm(header_h)
            hdr.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
            cell = hdr.cells[0].merge(hdr.cells[1]).merge(hdr.cells[2]).merge(hdr.cells[3])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            clean_para(p, WD_ALIGN_PARAGRAPH.CENTER)
            set_run(p.add_run(f"单词速记卡（第{day_no}天：{idx + 1:03d}/{total:03d}）"),
                    15, bold=True, color=(0x0D, 0x37, 0x6E))
            if title:
                sub = cell.add_paragraph()
                clean_para(sub, WD_ALIGN_PARAGRAPH.CENTER)
                set_run(sub.add_run(title), 9, color=(0x88, 0x88, 0x88))

            row_mm = (card_h - header_h - 2.0) / len(card)
            row_h = Mm(row_mm)
            en_wrap_w = en_w_pt - 1.5
            cn_wrap_w = cn_w_pt - 1.5
            en_size = _pick_card_font_size([e["text"] for e in card], bold_font,
                                           en_wrap_w, CARD_EN_FONT, row_mm * MM_TO_PT)
            cn_size = _pick_card_font_size([e["translation"] for e in card], font,
                                           cn_wrap_w, CARD_CN_FONT, row_mm * MM_TO_PT)

            for r, entry in enumerate(card):
                row = table.rows[r + 1]
                row.height = row_h
                row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

                texts = [
                    (str(r + 1), 10, False, (0xAA, 0xAA, 0xAA), None),
                    (entry["text"], en_size, True, (0x1A, 0x1A, 0x1A),
                     (bold_font, en_wrap_w)),
                    (entry["pos"], 10, False, (0x88, 0x88, 0x88), None),
                    (entry["translation"], cn_size, False, (0x15, 0x65, 0xC0),
                     (font, cn_wrap_w)),
                ]
                for col_i, (text, size, bold, color, wrap) in enumerate(texts):
                    item = row.cells[col_i]
                    item.width = col_w[col_i]
                    item.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    para = item.paragraphs[0]
                    clean_para(para)
                    para.paragraph_format.line_spacing = Pt(size * LINE_LEADING)
                    if not text:
                        continue
                    lines = _wrap_lines(text, wrap[0], size, wrap[1]) if wrap else [text]
                    for k, line in enumerate(lines):
                        run = para.add_run(line)
                        set_run(run, size, bold=bold, color=color)
                        if k < len(lines) - 1:
                            run.add_break()

            for col_i, width in enumerate(col_w):
                table.columns[col_i].width = width

            is_last = idx == total - 1
            if not is_last and idx % 2 == 0:
                spacer(Mm(4.0), dashed_line=True)
            elif is_last:
                spacer(Pt(1))
            else:
                spacer(Pt(1), page_break=True)

        doc.save(out_path)
        return out_path

    # ---------- 测验生成 ----------

    def generate_quiz(self, day_no, count=10, mode="mixed"):
        """生成选择题测验。
        mode: en2cn（英选中）/ cn2en（中选英）/ mixed（混合）
        返回 [{direction, question, question_extra, options, answer_index, entry}]
        """
        entries = [e for e in self.get_day_entries(day_no) if e["translation"]]
        if not entries:
            return []
        rng = random.Random()
        order = entries[:]
        rng.shuffle(order)
        picked = order[:count] if count < len(order) else order[:]

        cn_pool = []
        for e in entries:
            t = e["translation"]
            if t not in cn_pool:
                cn_pool.append(t)

        questions = []
        for e in picked:
            if mode == "mixed":
                direction = rng.choice(["en2cn", "cn2en"])
            else:
                direction = mode

            if direction == "en2cn":
                correct = e["translation"]
                distractors = [t for t in cn_pool if t != correct]
                rng.shuffle(distractors)
                options = [correct] + distractors[:3]
                question = e["text"]
                extra = e["pos"]
            else:
                correct = e["text"]
                pool = [x["text"] for x in entries if x["text"] != correct]
                unique = []
                for t in pool:
                    if t not in unique:
                        unique.append(t)
                rng.shuffle(unique)
                options = [correct] + unique[:3]
                question = e["translation"]
                extra = ""

            if len(options) < 2:
                continue
            rng.shuffle(options)
            questions.append({
                "direction": direction,
                "question": question,
                "question_extra": extra,
                "options": options,
                "answer_index": options.index(correct),
                "entry": e,
            })
        return questions


# ---------- 字体与文本度量辅助（PDF / DOCX 共用） ----------

_PDF_FONTS = None
_PDF_METRICS = None

MM_TO_PT = 72 / 25.4
LINE_LEADING = 1.3
MAX_TEXT_LINES = 3
_NO_LINE_START = set("。，、！？：；）》」』】.%")


def _string_width(text, font_name, size):
    global _PDF_METRICS
    if _PDF_METRICS is None:
        try:
            from reportlab.pdfbase import pdfmetrics
        except ImportError:
            raise RuntimeError("缺少 reportlab 库，请先安装：pip install reportlab")
        _PDF_METRICS = pdfmetrics
    return _PDF_METRICS.stringWidth(text, font_name, size)


def _ensure_pdf_fonts():
    """注册中文字体（微软雅黑优先），返回 (regular, bold) 字体名"""
    global _PDF_FONTS
    if _PDF_FONTS:
        return _PDF_FONTS
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError:
        raise RuntimeError("缺少 reportlab 库，请先安装：pip install reportlab")

    regular = bold = None
    candidates = [
        ("MSYH", "C:/Windows/Fonts/msyh.ttc", "MSYH-BD", "C:/Windows/Fonts/msyhbd.ttc"),
        ("SIMHEI", "C:/Windows/Fonts/simhei.ttf", "SIMHEI", "C:/Windows/Fonts/simhei.ttf"),
        ("SIMSUN", "C:/Windows/Fonts/simsun.ttc", "SIMSUN", "C:/Windows/Fonts/simsun.ttc"),
    ]
    for reg_name, reg_path, bold_name, bold_path in candidates:
        if os.path.exists(reg_path):
            try:
                pdfmetrics.registerFont(TTFont(reg_name, reg_path, subfontIndex=0))
                regular = reg_name
                if os.path.exists(bold_path):
                    try:
                        pdfmetrics.registerFont(TTFont(bold_name, bold_path, subfontIndex=0))
                        bold = bold_name
                    except Exception:
                        bold = None
                break
            except Exception:
                continue
    if not regular:
        regular = bold = "Helvetica"
    if not bold:
        bold = regular
    _PDF_FONTS = (regular, bold)
    return _PDF_FONTS


def _is_cjk(ch):
    o = ord(ch)
    return (0x2E80 <= o <= 0x9FFF or 0x3000 <= o <= 0x303F or 0xFF00 <= o <= 0xFFEF
            or 0xF900 <= o <= 0xFAFF or 0x20000 <= o <= 0x2FA1F)


def _tokenize(text):
    """切分为不可再拆的单位：单个中日韩字符 / 一段拉丁词 / None（表示原文此处有空格）。"""
    units = []
    buf = ""
    for ch in text:
        if ch.isspace():
            if buf:
                units.append(buf)
                buf = ""
            units.append(None)
        elif _is_cjk(ch):
            if buf:
                units.append(buf)
                buf = ""
            units.append(ch)
        else:
            buf += ch
    if buf:
        units.append(buf)
    return units


def _wrap_lines(text, font_name, size, max_width_pt):
    """折行：英文单词与中文单字各自为不可拆单位，句点等标点不单独落行首。"""
    if not text:
        return []
    units = []
    for unit in _tokenize(text):
        if unit is None:
            if units and units[-1] is not None:
                units.append(None)
            continue
        if units and units[-1] is not None and unit in _NO_LINE_START:
            units[-1] += unit
        else:
            units.append(unit)

    lines = []
    cur = ""
    pending_space = False
    for unit in units:
        if unit is None:
            if cur:
                pending_space = True
            continue
        cand = cur + (" " if pending_space else "") + unit
        if not cur or _string_width(cand, font_name, size) <= max_width_pt:
            cur = cand
            pending_space = False
        else:
            lines.append(cur)
            cur = unit
            pending_space = False
    if cur:
        lines.append(cur)
    return lines


def _pos_col_width(pos_list, font_name, size_pt, min_pt=10 * MM_TO_PT,
                   max_pt=26 * MM_TO_PT):
    """词性列宽：需容纳当天最长的词性标注。"""
    need = max([_string_width(p, font_name, size_pt) for p in pos_list if p] or [0.0])
    return min(max(need + 1 * MM_TO_PT, min_pt), max_pt)


def _split_card_columns(en_texts, cn_texts, en_font, cn_font, en_base, cn_base, avail_pt):
    """按两列最长词条约两行所需的宽度比例分配英/中列宽，返回 (en_w, cn_w)。"""
    en_need = max([_string_width(t, en_font, en_base) for t in en_texts] or [1.0]) / 2
    cn_need = max([_string_width(t, cn_font, cn_base) for t in cn_texts] or [1.0]) / 2
    ratio = en_need / max(en_need + cn_need, 1.0)
    ratio = min(max(ratio, 0.34), 0.56)
    return avail_pt * ratio, avail_pt * (1 - ratio)


def _pick_card_font_size(texts, font_name, max_width_pt, start_size, row_height_pt,
                         min_size=8.5):
    """整张卡使用统一字号：取能容纳最长词条的最大字号。"""
    size = start_size
    while size > min_size:
        n_lines = max((len(_wrap_lines(t, font_name, size, max_width_pt)) for t in texts),
                      default=1)
        if n_lines <= MAX_TEXT_LINES and n_lines * size * LINE_LEADING <= row_height_pt:
            return size
        size -= 0.5
    return min_size


def _draw_lines(canvas_obj, font_name, text, x, max_width_pt, center_y, size, rgb):
    """在行框内垂直居中绘制（可能多行的）文本。"""
    lines = _wrap_lines(text, font_name, size, max_width_pt)
    if not lines:
        return
    leading = size * LINE_LEADING
    canvas_obj.setFillColorRGB(*rgb)
    canvas_obj.setFont(font_name, size)
    y = center_y + (len(lines) - 1) * leading / 2 - size * 0.36
    for line in lines:
        canvas_obj.drawString(x, y, line)
        y -= leading
